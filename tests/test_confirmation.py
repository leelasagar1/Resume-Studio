import asyncio
import io
import pytest
from docx import Document
from app.confirmation import finalize_pending, without_terms
from app.demo import DemoProvider, SAMPLE_JOB, SAMPLE_RESUME
from app.models import GenerateRequest, Resume
from app.workflow import run_workflow


async def silent(*_):
    pass


@pytest.fixture
def result():
    request = GenerateRequest(resume_text=SAMPLE_RESUME, job_text=SAMPLE_JOB, confirmed=True, max_revisions=0)
    return asyncio.run(run_workflow(request, DemoProvider(delay=0), silent))


def test_pending_run_reports_both_bounds(result):
    assert result['report']['score'] == 100
    assert result['floor_score'] < 80
    assert set(result['unverified_terms']) == {'Dashboards', 'Data quality', 'Docker', 'Airflow'}


def test_rejected_terms_and_bullets_never_reach_word(result):
    assert [p['id'] for p in result['proposals']] == ['j1p1', 'e2p2']
    final = finalize_pending(result['_pending'], [], [])
    text = '\n'.join(p.text for p in Document(io.BytesIO(final['docx'])).paragraphs)
    assert 'Docker' not in text and 'Airflow' not in text and 'Containerized' not in text and 'SQL' in text
    assert final['report']['score'] == result['floor_score']
    assert final['confirmation_decisions'] == [{'term': 'Dashboards', 'accepted': False}, {'term': 'Data quality', 'accepted': False},
                                               {'term': 'Docker', 'accepted': False}, {'term': 'Airflow', 'accepted': False},
                                               {'term': result['proposals'][0]['text'], 'accepted': False},
                                               {'term': result['proposals'][1]['text'], 'accepted': False}]


def test_accepted_bullet_is_kept_as_confirmed_fact(result):
    final = finalize_pending(result['_pending'], [], ['j1p1'])
    text = '\n'.join(p.text for p in Document(io.BytesIO(final['docx'])).paragraphs)
    assert 'Containerized recurring Python reporting jobs with Docker' in text
    bullet = [b for b in Resume.model_validate(final['resume']).sections[0].entries[0].bullets if b.id == 'j1p1'][0]
    assert bullet.proposed is False and bullet.evidence_ids == ['U2'] and 'U2' in final['evidence']
    assert final['report']['score'] > result['floor_score']
    with pytest.raises(ValueError, match='unknown or duplicate proposed'):
        finalize_pending(result['_pending'], [], ['j1b1'])


def test_partial_acceptance_recomputes_score(result):
    final = finalize_pending(result['_pending'], ['docker'], [])  # case-insensitive match
    text = '\n'.join(p.text for p in Document(io.BytesIO(final['docx'])).paragraphs)
    assert 'Docker' in text and 'Airflow' not in text
    assert result['floor_score'] < final['report']['score'] < 100
    assert final['report']['unverified_terms'] == []


@pytest.mark.parametrize('terms', [['Kubernetes'], ['Docker', 'Docker'], ['SQL']])
def test_confirmation_cannot_inject_or_duplicate(result, terms):
    with pytest.raises(ValueError, match='unknown or duplicate'):
        finalize_pending(result['_pending'], terms, [])


def test_without_terms_keeps_evidenced_items_with_same_text():
    from app.demo import sample_draft
    draft = sample_draft(1)
    draft.sections[1].skill_groups[0].items[0].text = 'Docker'  # evidenced item that happens to share the term
    kept = without_terms(draft, {'Docker'})
    items = [i.text for g in kept.sections[1].skill_groups for i in g.items]
    assert items.count('Docker') == 1 and 'Airflow' in items
