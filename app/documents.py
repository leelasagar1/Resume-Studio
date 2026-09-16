import io
import re
import zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from .models import Resume, claims, proposed_claims, skill_items

MAX_FILE = 5 * 1024 * 1024


def extract_text(filename: str, data: bytes) -> str:
    if not data or len(data) > MAX_FILE:
        raise ValueError('Choose a nonempty file smaller than 5 MB.')
    ext = Path(filename).suffix.lower()
    try:
        if ext == '.txt':
            text = data.decode('utf-8-sig')
        elif ext == '.docx':
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 25 * 1024 * 1024:
                    raise ValueError('This Word file expands beyond the 25 MB limit.')
                if len(archive.infolist()) > 2000:
                    raise ValueError('This Word file has too many embedded parts.')
                if 'word/document.xml' not in archive.namelist():
                    raise ValueError('This is not a valid DOCX file.')
            doc = Document(io.BytesIO(data))
            # Include tables, headers and footers often used in existing resumes.
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            parts = []
            for section in doc.sections:
                parts.extend(p.text for p in section.header.paragraphs)
            for block in doc.iter_inner_content():
                if isinstance(block, Paragraph):
                    parts.append(block.text)
                elif isinstance(block, Table):
                    parts.extend(' | '.join(c.text for c in row.cells) for row in block.rows)
            for section in doc.sections:
                parts.extend(p.text for p in section.footer.paragraphs)
            text = '\n'.join(parts)
        elif ext == '.pdf':
            if not data.startswith(b'%PDF-'):
                raise ValueError('This is not a valid PDF.')
            pdf = PdfReader(io.BytesIO(data))
            if pdf.is_encrypted:
                raise ValueError('Upload an unencrypted PDF or a DOCX file.')
            if len(pdf.pages) > 12:
                raise ValueError('Use a PDF of at most 12 pages.')
            text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
        else:
            raise ValueError('Supported formats: PDF, DOCX, and UTF-8 TXT.')
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('The file could not be read. Try DOCX or paste its text.') from exc
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text).strip()
    if len(text) < 40:
        raise ValueError('Not enough text could be extracted. For scanned PDFs, paste the text or upload DOCX.')
    if len(text) > 20000:
        raise ValueError('The file contains too much text. Use a shorter document or paste the relevant text.')
    return text


def render_docx(resume: Resume) -> tuple[bytes, dict]:
    if any(not item.evidence_ids for item in skill_items(resume)) or proposed_claims(resume):
        raise ValueError('Confirm or remove proposed content before generating a submission-ready document.')
    doc = Document()
    # Some bundled Word defaults include a blue Title border. Keep the resume plain.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.48)
    sec.left_margin = sec.right_margin = Inches(.62)
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(9.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0
    for name in ['Title', 'Heading 1', 'Heading 2']:
        doc.styles[name].font.name = 'Calibri'
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Title'].font.size = Pt(21)
    doc.styles['Title'].paragraph_format.space_after = Pt(4)
    doc.styles['Heading 1'].font.size = Pt(11)
    doc.styles['Heading 1'].font.bold = True
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(6)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(2)
    doc.add_paragraph(resume.name.text, 'Title')
    if resume.headline.strip():
        head = doc.add_paragraph()
        head.add_run(resume.headline.strip()).bold = True
    if resume.contact:
        doc.add_paragraph(' | '.join(c.text for c in resume.contact))
    if resume.summary:
        doc.add_heading('Summary', level=1)
        for claim in resume.summary:
            doc.add_paragraph(claim.text)
    for section in resume.sections:
        doc.add_heading(section.heading, level=1)
        for entry in section.entries:
            p = doc.add_paragraph()
            p.add_run(entry.heading.text).bold = True
            p.paragraph_format.keep_with_next = bool(entry.detail or entry.bullets)
            if entry.detail:
                p = doc.add_paragraph(entry.detail.text)
                p.paragraph_format.keep_with_next = bool(entry.bullets)
            for bullet in entry.bullets:
                p = doc.add_paragraph(bullet.text, 'List Bullet')
                p.paragraph_format.space_after = Pt(1)
        for group in section.skill_groups:
            p = doc.add_paragraph()
            p.add_run(group.label + ': ').bold = True
            p.add_run(', '.join(item.text for item in group.items))
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.title = 'Resume'
    for border in doc.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    buffer = io.BytesIO()
    doc.save(buffer)
    data = buffer.getvalue()
    extracted = '\n'.join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    expected = [c.text for c in claims(resume)] + [i.text for i in skill_items(resume)] + [resume.headline.strip()]
    missing = [text for text in expected if text and text not in extracted]
    if missing:
        raise ValueError('Document round-trip check failed.')
    return data, {'text_round_trip': True, 'single_column': True,
                  'visual_check': 'Open the downloaded file to confirm page breaks before applying.'}
